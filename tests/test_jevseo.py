"""Offline tests: no network, no Jev spend. Run with `python3 -m unittest discover -s tests -v`."""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from jevseo import checks, jev, score  # noqa: E402
from jevseo.parse import normalize_url, parse_page  # noqa: E402

HOME = "https://example.org/"
PAGE_HTML = """<!doctype html><html lang="en"><head>
<title>Acme Plumbing | Emergency plumbers in Leeds</title>
<meta name="description" content="24 hour emergency plumbing across Leeds.">
<meta name="viewport" content="width=device-width">
<link rel="canonical" href="https://example.org/services/emergency/">
<script type="application/ld+json">{"@context":"https://schema.org","@type":"Plumber","name":"Acme"}</script>
<script type="application/ld+json">{ broken json </script>
<link rel="stylesheet" href="http://cdn.example.org/site.css">
</head><body><nav><a href="/">Home</a><a href="/about">About</a></nav>
<main><h1>Emergency plumbers</h1><h3>Skipped level</h3>
<p>We fix burst pipes across Leeds within 60 minutes, day or night.</p>
<a href="http://other.org/plain-link">an outbound http link is not mixed content</a>
<a href="/contact">click here</a><img src="/a.jpg"><img src="/b.jpg" alt="" width="10" height="10">
</main></body></html>"""


def page(url, **kw):
    base = parse_page(url, PAGE_HTML, "example.org")
    base.update({"url": url, "final_url": url, "kind": "page", "status": 200, "depth": 1, "inlinks": 2, "in_sitemap": True, "ttfb_ms": 120, "bytes": 5000, "headers": {}, "rendered": False, "js_dependent": False})
    base.update(kw)
    return base


def crawl_fixture():
    home = page(HOME, depth=0, canonical=HOME, title="Acme Plumbing", favicon=True, schema_types=["WebSite"])
    svc = page("https://example.org/services/emergency", canonical="https://example.org/services/emergency")
    dup = page("https://example.org/services/emergency-2", canonical="https://example.org/services/emergency-2", in_sitemap=False, inlinks=0)
    thin = page("https://example.org/thin", canonical="https://example.org/thin", word_count=40, meta_robots="noindex")
    moved = {"url": "https://example.org/old", "kind": "redirect", "status": 301, "final_url": HOME, "redirect_chain": [{"url": "https://example.org/old", "status": 301}, {"url": "https://example.org/mid", "status": 302}], "depth": 1}
    gone = {"url": "https://example.org/gone", "kind": "page", "status": 404, "depth": 2}
    home["links_internal"] = [{"url": "https://example.org/gone", "anchor": "read more"}, {"url": "https://example.org/old", "anchor": "Old"}]
    return {
        "input_url": HOME, "start_url": HOME, "final_url": HOME, "origin": "https://example.org", "domain": "example.org", "https": True,
        "robots": {"url": "https://example.org/robots.txt", "status": 200, "present": True, "bytes": 30, "sitemaps": [], "crawl_delay": None, "disallow_all": False,
                   "search_bots": {"Googlebot": True, "Bingbot": True}, "ai_bots": {"GPTBot": False, "ClaudeBot": True}, "ai_bots_named": ["GPTBot"]},
        "sitemaps": {"files": [{"url": "https://example.org/sitemap.xml", "status": 200, "urls": 4, "error": None, "lastmod": 0}], "urls": [HOME, "https://example.org/services/emergency", "https://example.org/thin", "https://example.org/old"], "total_urls": 4},
        "probes": {"http_to_https": True, "host_variant": {"host": "www.example.org", "status": 200, "redirects_to_canonical_host": False}, "not_found_status": 200, "llms_txt": False},
        "pages": [home, svc, dup, thin, moved, gone],
        "link_status": {"https://example.org/gone": 404, "https://example.org/old": 301, "https://example.org/contact": 200},
        "external_status": {"http://other.org/plain-link": 200}, "external_total": 1, "blocked_by_robots": [],
        "render": {"mode": "auto", "browser_available": None, "homepage_rendered": False}, "requests": 12,
        "limits": {"max_pages": 60, "max_depth": 5, "time_budget_s": 600, "hit_page_cap": False, "hit_time_budget": False},
    }


def answer(qtype, value, conf=0.9):
    if qtype == "noul":
        return {"type": "noul", "value": value, "band": "yes" if value >= 0.8 else "no" if value <= 0.2 else "review"}
    band = "act" if conf >= jev.ACT else "review"
    if qtype == "choice":
        return {"type": "choice", "value": value, "probabilities": {value: conf}, "confidence": conf, "band": band}
    return {"type": "score", "value": value, "raw": value * 3, "levels": 4, "probabilities": {"0": 0.1, "3": 0.9}, "confidence": conf, "band": band}


def judged_fixture(crawl):
    pages = {}
    for p in checks.html_pages(crawl):
        pages[p["url"]] = {
            "page_type": answer("choice", "product_or_service"), "intent": answer("choice", "local"), "importance": answer("score", 1.0),
            "action": answer("choice", "rewrite", 0.6), "helpfulness": answer("score", 0.2), "specificity": answer("score", 0.3), "trust": answer("score", 0.8),
            "answer_first": answer("noul", 0.1), "citable": answer("score", 0.3, 0.5), "clear_next_step": answer("noul", 0.9), "title_fit": answer("score", 0.9), "meta_fit": answer("score", 0.9), "h1_fit": answer("noul", 0.95),
        }
    site = {"business_model": answer("choice", "local_service"), "value_prop": answer("score", 0.4), "entity_clarity": answer("noul", 0.1),
            "topical_focus": answer("score", 0.9), "serves_local_area": answer("noul", 0.95)}
    return {"available": True, "site": site, "pages": pages, "pairs": [{"a": HOME, "b": "https://example.org/services/emergency", "title_overlap": 0.5, "judgment": answer("noul", 0.9)}],
            "ledger": {"model_requested": "jev-latest", "model_returned": "jev-1.13.0", "requests": 5, "failed": 0, "input_tokens": 12000, "output_tokens": 900, "skipped_budget": 0, "errors": [], "cost_usd": 0.000504, "usd_per_mtok": 0.042},
            "questions": {"site": jev.site_questions()}}


class ParseTests(unittest.TestCase):
    def test_facts(self):
        p = parse_page("https://example.org/x", PAGE_HTML, "example.org")
        self.assertEqual(p["title"], "Acme Plumbing | Emergency plumbers in Leeds")
        self.assertEqual(p["canonical"], "https://example.org/services/emergency")
        self.assertIn("Plumber", p["schema_types"])
        self.assertEqual(len(p["schema_errors"]), 1)
        self.assertEqual(p["heading_skips"], 1)
        self.assertEqual(p["images"]["missing_alt"], 1)
        self.assertEqual(p["images"]["empty_alt"], 1)
        self.assertEqual(p["lang"], "en")

    def test_mixed_content_ignores_plain_links(self):
        p = parse_page("https://example.org/x", PAGE_HTML, "example.org")
        self.assertEqual(p["mixed_content"], ["http://cdn.example.org/site.css"])

    def test_normalize(self):
        self.assertEqual(normalize_url("HTTPS://Example.org/a/#top"), "https://example.org/a")
        self.assertEqual(normalize_url("https://example.org"), "https://example.org/")


class CheckTests(unittest.TestCase):
    def setUp(self):
        self.crawl = crawl_fixture()
        self.found = {f["id"]: f for f in checks.run_checks(self.crawl)}

    def test_expected_rules_fire(self):
        for rule in ["http_errors", "broken_internal_links", "links_to_redirects", "redirect_chains", "noindex", "soft_404", "host_variant",
                     "jsonld_errors", "mixed_content", "images_alt", "heading_skips", "thin_content", "generic_anchors", "ai_bots_blocked", "llms_txt_missing", "sitemap_bad_urls"]:
            self.assertIn(rule, self.found, rule)

    def test_structured_data_required_properties(self):
        crawl = crawl_fixture()
        crawl["pages"][1]["schema_nodes"] = [
            {"types": ["SoftwareApplication"], "keys": ["name", "offers"], "offers_price": True, "list_items": 0, "list_items_ok": False},
            {"types": ["Product"], "keys": ["name", "offers"], "offers_price": True, "list_items": 0, "list_items_ok": False},
        ]
        crawl["pages"][1]["schema_types"] = ["FAQPage", "Product", "SoftwareApplication"]
        found = {f["id"]: f for f in checks.run_checks(crawl)}
        self.assertIn("aggregateRating or review", found["schema_required"]["evidence"])
        self.assertNotIn("Product on", found["schema_required"]["evidence"])  # complete Product is not flagged
        self.assertEqual(found["faq_rich_result_limited"]["severity"], "info")

    def test_heuristics_are_labelled(self):
        self.assertTrue(self.found["thin_content"]["heuristic"])
        self.assertFalse(self.found["http_errors"]["heuristic"])

    def test_every_finding_has_a_source(self):
        for f in self.found.values():
            self.assertTrue(f["source"].startswith("https://"), f["id"])

    def test_no_https_not_raised_for_https_site(self):
        self.assertNotIn("no_https", self.found)


class JevValidationTests(unittest.TestCase):
    def test_score_normalised(self):
        q = jev.score("x", ["a", "b", "c", "d"])
        a = jev.validate(q, {"score": 1.5, "confidence": 0.9, "probabilities": {}})
        self.assertEqual(a["value"], 0.5)
        self.assertEqual(a["band"], "act")

    def test_noul_bands(self):
        q = jev.noul("x", "t", "f")
        self.assertEqual(jev.validate(q, {"noul": 0.5})["band"], "review")
        self.assertEqual(jev.validate(q, {"noul": 0.1})["band"], "no")

    def test_rejects_invalid(self):
        with self.assertRaises(RuntimeError):
            jev.validate(jev.choice("x", {"a": None}), {"choice": "b", "confidence": 1})
        with self.assertRaises(RuntimeError):
            jev.validate(jev.score("x", ["a", "b"]), {"score": 3, "confidence": 1})

    def test_optional_questions_skipped_when_absent(self):
        q = jev.page_questions({"title": None, "meta_description": None, "h1": []})
        self.assertNotIn("meta_fit", q)
        self.assertNotIn("title_fit", q)
        self.assertNotIn("h1_fit", q)

    def test_budget_cap_skips_without_network(self):
        client = jev.Jev(budget_usd=0.0, log=lambda *_: None)
        client.key = "test-key-not-used"
        self.assertIsNone(client.ask({"x": "y" * 1000}, {"q": jev.noul("x", "t", "f")}))
        self.assertEqual(client.ledger["skipped_budget"], 1)
        self.assertEqual(client.ledger["requests"], 0)

    def test_pair_shortlist(self):
        a = {"url": "a", "title": "Emergency plumber Leeds", "h1": ["Emergency plumber"], "text_hash": "1"}
        b = {"url": "b", "title": "Emergency plumber Leeds night", "h1": [], "text_hash": "2"}
        c = {"url": "c", "title": "Bathroom fitting", "h1": [], "text_hash": "3"}
        pairs = jev.overlap_candidates([a, b, c])
        self.assertEqual([(x["url"], y["url"]) for x, y, _ in pairs], [("a", "b")])


class ScoreTests(unittest.TestCase):
    def setUp(self):
        self.crawl = crawl_fixture()
        self.judged = judged_fixture(self.crawl)
        self.findings = checks.run_checks(self.crawl) + score.jev_findings(self.crawl, self.judged)

    def test_jev_findings(self):
        ids = {f["id"] for f in self.findings}
        for rule in ["jev_value_prop", "jev_entity_clarity", "jev_local_schema", "jev_helpfulness", "jev_answer_first", "jev_rewrite", "jev_cannibalization"]:
            self.assertIn(rule, ids, rule)
        self.assertNotIn("jev_trust", ids)
        self.crawl["pages"][0]["schema_types"] = ["Plumber"]
        self.assertNotIn("jev_local_schema", {f["id"] for f in score.jev_findings(self.crawl, self.judged)})
        citable = next(f for f in self.findings if f["id"] == "jev_citable")
        self.assertEqual(citable["needs_review"], citable["count"])

    def test_scores_and_missing_areas(self):
        s = score.score(self.crawl, self.findings, self.judged, None)
        self.assertTrue(0 <= s["overall"] <= 100)
        no_jev = score.score(self.crawl, checks.run_checks(self.crawl), {"pages": {}}, None)
        self.assertIsNone(no_jev["categories"]["content"])
        self.assertEqual(no_jev["completeness"]["categories_scored"], 7)

    def test_blocked_site_is_capped(self):
        crawl = crawl_fixture()
        crawl["robots"]["disallow_all"] = True
        s = score.score(crawl, checks.run_checks(crawl), {"pages": {}}, None)
        self.assertLessEqual(s["overall"], 20)

    def test_actions_ranked_and_ided(self):
        acts = score.actions(self.findings, self.judged, 4)
        self.assertEqual(acts[0]["action_id"], "JEV-001")
        order = [a["priority"] for a in acts]
        self.assertEqual(order, sorted(order))
        self.assertTrue(all(0 <= a["impact"] <= 100 for a in acts))


class ReviewRegressionTests(unittest.TestCase):
    def test_sitemap_with_comments_parses(self):
        from unittest import mock
        from jevseo.crawl import Fetcher, read_sitemaps

        xml = b'<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><!-- Main --><url><loc>https://x.org/</loc><lastmod>2026-09-10</lastmod></url><url><loc>https://x.org/a</loc></url></urlset>'

        class R:
            status_code = 200
            content = xml

        with mock.patch.object(Fetcher, "get", return_value=R()):
            out = read_sitemaps(Fetcher(), ["https://x.org/sitemap.xml"])
        self.assertEqual(out["total_urls"], 2)
        self.assertEqual(out["files"][0]["lastmod"], 1)

    def test_utf8_without_header_charset_is_not_mojibake(self):
        from jevseo.crawl import decode

        class R:
            headers = {"content-type": "text/html"}
            content = "<html><head><meta charset='utf-8'></head><body>©2026 café ✢</body></html>".encode("utf-8")
            text = content.decode("latin-1")

        self.assertIn("©2026 café ✢", decode(R()))

    def test_noise_tokens_removed_from_text(self):
        p = parse_page("https://x.org/", "<html><body><main><p>Hello " + "QUJD" * 30 + " world</p></main></body></html>", "x.org")
        self.assertEqual(p["text_excerpt"], "Hello world")

    def test_unicode_urls_normalise_like_requests(self):
        self.assertEqual(normalize_url("https://x.org/café/"), normalize_url("https://x.org/caf%C3%A9"))

    def test_homepage_hreflang_self_reference_passes(self):
        crawl = crawl_fixture()
        crawl["pages"][0]["hreflang"] = [{"lang": "en", "href": "https://example.org"}, {"lang": "x-default", "href": "https://example.org/"}]
        self.assertNotIn("hreflang_issues", {f["id"] for f in checks.run_checks(crawl)})

    def test_robots_blocked_target_is_not_broken(self):
        crawl = crawl_fixture()
        crawl["pages"][0]["links_internal"] = [{"url": "https://example.org/private", "anchor": "Private"}]
        crawl["link_status"] = {"https://example.org/private": "blocked"}
        self.assertNotIn("broken_internal_links", {f["id"] for f in checks.run_checks(crawl)})

    def test_link_through_redirect_to_404_is_broken(self):
        crawl = crawl_fixture()
        crawl["pages"][0]["links_internal"] = [{"url": "https://example.org/moved", "anchor": "Moved"}]
        crawl["link_status"] = {"https://example.org/moved": 404}
        self.assertIn("broken_internal_links", {f["id"] for f in checks.run_checks(crawl)})

    def test_malformed_jev_response_is_recorded_not_raised(self):
        from unittest import mock

        class Resp:
            status_code = 200
            headers = {}
            text = ""

            def __init__(self, payload):
                self.payload = payload

            def json(self):
                if self.payload is None:
                    raise ValueError("not json")
                return self.payload

        q = {"a": jev.noul("x", "t", "f")}
        for payload in (None, {"answers": {"a": {}}}, {"answers": {"a": {"noul": None}}, "usage": None}, {"answers": {}}):
            client = jev.Jev(budget_usd=1.0, log=lambda *_: None)
            client.key = "test-key-not-used"
            with mock.patch("jevseo.jev.requests.post", return_value=Resp(payload)):
                self.assertIsNone(client.ask({"s": 1}, q))
            self.assertEqual(client.ledger["failed"], 1)

    def test_missing_usage_still_counts_toward_budget(self):
        from unittest import mock

        class Resp:
            status_code = 200
            headers = {}
            text = ""

            def json(self):
                return {"answers": {"a": {"noul": 0.9}}, "model": "jev-1.13.0"}

        client = jev.Jev(budget_usd=1.0, log=lambda *_: None)
        client.key = "test-key-not-used"
        with mock.patch("jevseo.jev.requests.post", return_value=Resp()):
            self.assertEqual(client.ask({"s": "x" * 300}, {"a": jev.noul("x", "t", "f")})["a"]["band"], "yes")
        self.assertGreater(client.ledger["input_tokens"], 0)
        self.assertEqual(client.ledger["usage_estimated"], 1)


def dfs_fixture():
    return {
        "available": True, "location_code": 2840, "language_code": "en",
        "overview": {"count": 3, "etv": 120.5, "pos_1": 0, "pos_2_3": 1, "pos_4_10": 1, "pos_11_20": 1},
        "ranked": [{"keyword": "emergency plumber leeds", "position": 7, "volume": 900, "difficulty": 20, "intent": "commercial", "url": "https://example.org/services/emergency"},
                   {"keyword": "acme plumbing", "position": 2, "volume": 50, "difficulty": 5, "intent": "navigational", "url": HOME}],
        "ranked_total": 2,
        "opportunities": [{"keyword": "burst pipe repair", "volume": 500, "difficulty": 10, "intent": "transactional", "source": "idea"},
                          {"keyword": "burst-pipe repair", "volume": 400, "difficulty": 10, "intent": "transactional", "source": "idea"},
                          {"keyword": "boiler service leeds", "volume": 300, "difficulty": 30, "intent": "commercial", "source": "idea"},
                          {"keyword": "british gas boiler", "volume": 9000, "difficulty": 60, "intent": "navigational", "source": "idea"},
                          {"keyword": "car insurance", "volume": 90000, "difficulty": 90, "intent": "commercial", "source": "idea"}],
        "competitors": [{"domain": "rival.co.uk", "shared_keywords": 2}], "competitors_all": [],
        "referring_domains": [{"domain": "example.org", "referring_domains": 3}, {"domain": "rival.co.uk", "referring_domains": 200}, {"domain": "other.co.uk", "referring_domains": 90}],
        "backlinks": {"backlinks": 10, "referring_domains": 3, "broken_backlinks": 2, "broken_pages": 1},
        "serps": [{"keyword": "emergency plumber leeds", "features": ["ai_overview"], "top": [], "own_position": 7, "ai_overview": True, "ai_overview_cites_site": False, "ai_overview_domains": ["rival.co.uk"]}],
        "mentions": {"total": 4, "items": [{"platform": "google", "model": "google_ai_overview", "question": "q", "ai_search_volume": 10, "sources": []}]},
        "ledger": {"requests": 17, "cost_usd": 0.3, "failed": 0, "skipped_budget": 0, "errors": [], "calls": []},
    }


def keyword_judgments():
    def kw(rel, page, other=0.1):
        return {"relevance": answer("score", rel), "page": answer("choice", page), "page_url": None if page == "none_fit" else "https://example.org/services/emergency", "other_brand": answer("noul", other)}
    return {"emergency plumber leeds": kw(1.0, "p1"), "acme plumbing": kw(1.0, "p0"), "burst pipe repair": kw(0.9, "p1"), "burst-pipe repair": kw(0.9, "p1"),
            "boiler service leeds": kw(0.8, "none_fit"), "british gas boiler": kw(0.9, "p1", other=0.95), "car insurance": kw(0.0, "none_fit")}


class DataForSEOTests(unittest.TestCase):
    def setUp(self):
        self.crawl = crawl_fixture()
        self.judged = judged_fixture(self.crawl) | {"keywords": keyword_judgments()}
        self.found = {f["id"]: f for f in score.dfs_findings(self.crawl, self.judged, dfs_fixture())}

    def test_findings(self):
        self.assertEqual(set(self.found), {"dfs_striking", "dfs_existing_page", "dfs_new_page", "dfs_backlink_gap", "dfs_broken_backlinks", "dfs_aio_not_cited"})

    def test_jev_filters_irrelevant_and_other_brands_and_duplicates(self):
        kept = [k["keyword"] for k in self.found["dfs_existing_page"]["detail"]["keywords"]] + [k["keyword"] for k in self.found["dfs_new_page"]["detail"]["keywords"]]
        self.assertNotIn("car insurance", kept)          # Jev relevance 0
        self.assertNotIn("british gas boiler", kept)     # another brand
        self.assertEqual(sum(1 for k in kept if "burst" in k), 1)  # spelling variants merged
        self.assertIn("boiler service leeds", kept)

    def test_visibility_scored_only_with_dataforseo(self):
        findings = checks.run_checks(self.crawl) + list(self.found.values())
        self.assertIsNotNone(score.score(self.crawl, findings, self.judged, None, dfs_fixture())["categories"]["visibility"])
        self.assertIsNone(score.score(self.crawl, findings, self.judged, None, None)["categories"]["visibility"])

    def test_budget_cap_blocks_before_dispatch(self):
        from unittest import mock
        from jevseo.dfs import DataForSEO

        client = DataForSEO(budget_usd=0.001, log=lambda *_: None)
        with mock.patch("jevseo.dfs.requests.post") as post:
            self.assertIsNone(client.post("backlinks/summary/live", {"target": "example.org"}))
            post.assert_not_called()
        self.assertEqual(client.ledger["skipped_budget"], 1)

    def test_unreported_cost_is_charged_at_estimate(self):
        from unittest import mock
        from jevseo.dfs import ESTIMATE, DataForSEO

        class R:
            status_code = 200

            def json(self):
                return {"tasks": [{"status_code": 20000, "result": [{"ok": 1}]}]}

        client = DataForSEO(budget_usd=1.0, log=lambda *_: None)
        with mock.patch("jevseo.dfs.requests.post", return_value=R()):
            self.assertEqual(client.post("backlinks/summary/live", {}), {"ok": 1})
        self.assertEqual(client.ledger["cost_usd"], ESTIMATE["backlinks/summary/live"])


class RenderTests(unittest.TestCase):
    def build(self, narrative=None, full=False):
        from jevseo.cli import digest
        from jevseo.report import build

        crawl = crawl_fixture()
        judged = judged_fixture(crawl) | ({"keywords": keyword_judgments()} if full else {})
        dfs = dfs_fixture() if full else None
        findings = checks.run_checks(crawl) + score.jev_findings(crawl, judged) + score.dfs_findings(crawl, judged, dfs)
        data = {"schema_version": "1.0", "tool": {"name": "jev-seo", "version": "test"},
                "run": {"started_at": "2026-09-21T10:00:00+00:00", "finished_at": "2026-09-21T10:05:00+00:00", "timings_s": {}, "options": {}},
                "site": {k: v for k, v in crawl.items() if k != "pages"}, "pages": crawl["pages"], "findings": findings,
                "passed_rules": checks.passed_rules(findings), "jev": judged, "performance": None, "dataforseo": dfs,
                "scores": score.score(crawl, findings, judged, None, dfs), "actions": score.actions(findings, judged, 4)}
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, True)
        (tmp / "audit.json").write_text(json.dumps(data))
        self.assertIn("JEV-001", digest(data))
        if narrative is not None:
            (tmp / "narrative.json").write_text(json.dumps(narrative))
        return tmp, build(tmp, ["pdf", "xlsx", "md"], log=lambda *_: None)

    def test_all_formats(self):
        tmp, out = self.build()
        self.assertEqual(out["pdf"].read_bytes()[:5], b"%PDF-")
        self.assertIn("## Priority actions", out["md"].read_text())
        from openpyxl import load_workbook

        wb = load_workbook(out["xlsx"])
        self.assertEqual(wb.sheetnames[:3], ["Summary", "Actions", "Pages"])
        ws = wb["Actions"]
        self.assertEqual(ws["E1"].value, "Status")
        self.assertEqual(ws["C1"].value, "Priority")
        summary = [c.value for row in wb["Summary"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(any("COUNTIF(Actions!$E:$E" in f for f in summary))
        self.assertTrue(ws.data_validations.dataValidation)

    def test_full_mode_renders_visibility(self):
        tmp, out = self.build(full=True)
        md = out["md"].read_text()
        self.assertIn("## Search visibility (DataForSEO)", md)
        self.assertIn("boiler service leeds", md)
        self.assertNotIn("car insurance", md)
        from openpyxl import load_workbook

        self.assertTrue({"Rankings", "Opportunities", "Competitors", "SERPs", "AI mentions"} <= set(load_workbook(out["xlsx"]).sheetnames))

    def test_narrative_rejects_unknown_ids(self):
        with self.assertRaises(SystemExit):
            self.build({"executive_summary": ["See JEV-999."], "strengths": [], "risks": [], "plan": []})

    def test_invented_numbers_flagged(self):
        tmp, out = self.build({"executive_summary": ["Score 57 with confidence 0.89 across JEV-001."], "strengths": [], "risks": [], "plan": []})
        from jevseo.report import load_narrative

        n = load_narrative(tmp, json.loads((tmp / "audit.json").read_text()))
        self.assertIn("0.89", n["unverified_numbers"])
        self.assertIn("57", n["unverified_numbers"])

    def test_narrative_used(self):
        tmp, out = self.build({"executive_summary": ["Custom verdict about JEV-001."], "strengths": ["s"], "risks": ["r"], "plan": [{"horizon": "This week", "items": ["JEV-001"]}]})
        self.assertIn("Custom verdict about JEV-001.", out["md"].read_text())


if __name__ == "__main__":
    unittest.main()
